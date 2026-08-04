#!/usr/bin/env python3
"""Build the Hyperion engineering handbook and its native diagrams."""

from __future__ import annotations

import argparse
import re
import textwrap
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import unquote, urlparse

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree
from PIL import Image, ImageDraw, ImageFont

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "17324D"
PALE_BLUE = "E8EEF5"
PALE_GOLD = "FFF4D6"
PALE_GREEN = "E8F4EC"
PALE_RED = "FBE9E7"
GREY = "66727F"
CONTENT_DXA = 9360
HANDBOOK_VERSION = "1.2"
HANDBOOK_DATE = "5 August 2026"
HANDBOOK_EXPECTED_PAGES = 49
HANDBOOK_RELEASED_UTC = datetime(2026, 8, 4, 18, 30, tzinfo=timezone.utc)
STATUS_STYLES = {
    "IMPLEMENTED": "2E74B5",
    "TARGET": "2B7A4B",
    "RESEARCH": "A66B00",
    "EXTERNAL / BLOCKED": "A33B30",
}
IMAGE_DESCRIPTIONS = {
    "assets/hyperion-system-architecture.png": (
        "Four-layer Hyperion architecture. Implemented client and coordination components send labelled "
        "requests, approvals, frozen specifications, and bound verification requests downward. Model-evidence "
        "and safety layers return quotes, inert artifacts, results, usage, and verifier evidence upward. Status "
        "badges distinguish implemented repository infrastructure from target, research, and externally blocked work."
    ),
    "assets/hyperion-run-state-machine.png": (
        "Hyperion run states Created, Running, Awaiting approval, Released, Cancelled, and Failed. Labelled arrows "
        "show start, candidate-ready, human-approval, cancellation or policy-stop, cleanup-fault, and retry or resume "
        "transitions; every transition is persisted with optimistic concurrency and hash-chained events."
    ),
    "assets/hyperion-evidence-lifecycle.png": (
        "Data, training, evaluation, and serving evidence layers connected by labelled evidence flows. Status badges "
        "show which contracts and harnesses exist and which production reviews, checkpoints, sealed tasks, or runtime "
        "attestation remain target or externally blocked."
    ),
    "assets/hyperion-trust-chain.png": (
        "Trust chain from evidence producers through independent review and detached signing to coordinator, hosted "
        "platform, and deployment consumers. Bidirectional arrows label evidence bundles, findings, canonical "
        "preimages, signatures, signed envelopes, and verification results; status badges expose missing production "
        "reviewers, signer custody, and KMS or HSM integration."
    ),
    "assets/hyperion-frontier-model-factory.png": (
        "Five-layer model factory covering authority, knowledge, teachers, student training, and runtime. Labelled "
        "flows carry policy gates, authorised examples, teacher traces, training specifications, candidates, metrics, "
        "accepted artifacts, telemetry, and rollback evidence. Status badges separate tested infrastructure from "
        "research, targets, and external blockers."
    ),
    "assets/hyperion-kimi-k3-adoption-map.png": (
        "Decision map from specific Kimi K3 report areas through Hyperion's evidence and feasibility gate. The map "
        "routes agent contracts and hidden verification to immediate adoption, quantisation and serving mechanisms to "
        "controlled evaluation, and backbone mechanisms to equal-compute pretraining research, while explicitly "
        "rejecting a direct full-model or QLoRA retrofit."
    ),
    "assets/hyperion-gpu-cluster-topology.png": (
        "Proposed small-cluster topology with control, storage, GPU, and isolation layers. Labelled arrows show signed "
        "run specifications, read-only snapshots, checkpoints, metrics, verification requests, verdicts, and recovery "
        "proof. Status badges show that immutable artifact contracts exist while the cluster, production corpus, KMS, "
        "GPU workers, and several isolation controls remain target or externally blocked."
    ),
}


def font(name: str, size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    fonts = Path("C:/Windows/Fonts")
    candidates = {
        "regular": ["segoeui.ttf", "arial.ttf"],
        "bold": ["seguisb.ttf", "segoeuib.ttf", "arialbd.ttf"],
    }
    for candidate in candidates["bold" if bold else "regular"]:
        path = fonts / candidate
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def wrapped(
    draw: ImageDraw.ImageDraw,
    value: str,
    box: tuple[int, int, int, int],
    value_font: ImageFont.FreeTypeFont,
    fill: str,
    spacing: int = 8,
) -> None:
    x1, y1, x2, y2 = box
    width = max(12, int((x2 - x1) / max(value_font.getlength("M"), 1)))
    lines: list[str] = []
    for paragraph in value.splitlines():
        lines.extend(
            textwrap.wrap(paragraph, width=width, break_long_words=False) or [""]
        )
    metrics = draw.multiline_textbbox(
        (0, 0), "\n".join(lines), font=value_font, spacing=spacing
    )
    height = metrics[3] - metrics[1]
    draw.multiline_text(
        ((x1 + x2) / 2, (y1 + y2 - height) / 2),
        "\n".join(lines),
        font=value_font,
        fill=fill,
        anchor="ma",
        align="center",
        spacing=spacing,
    )


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str = BLUE) -> None:
    if not color.startswith("#"):
        color = f"#{color}"
    draw.line([start, end], fill=color, width=6)
    x, y = end
    if abs(end[0] - start[0]) >= abs(end[1] - start[1]):
        direction = 1 if end[0] > start[0] else -1
        points = [(x, y), (x - 20 * direction, y - 13), (x - 20 * direction, y + 13)]
    else:
        direction = 1 if end[1] > start[1] else -1
        points = [(x, y), (x - 13, y - 20 * direction), (x + 13, y - 20 * direction)]
    draw.polygon(points, fill=color)


def layered_diagram(
    path: Path,
    title: str,
    subtitle: str,
    layers: list[tuple[str, list[tuple[str, str]], str]],
    edge_labels: list[tuple[str, str]],
) -> None:
    if len(edge_labels) != len(layers) - 1:
        raise ValueError(
            "Layered diagrams require one labelled edge pair between each layer"
        )
    image_height = max(1120, 370 + len(layers) * 238)
    image = Image.new("RGB", (1800, image_height), "#F7F9FB")
    draw = ImageDraw.Draw(image)
    title_font = font("regular", 52, True)
    subtitle_font = font("regular", 25)
    label_font = font("regular", 24, True)
    box_font = font("regular", 22, True)
    status_font = font("regular", 15, True)
    edge_font = font("regular", 18, True)
    draw.text((90, 68), title, font=title_font, fill="#17324D")
    draw.text((90, 136), subtitle, font=subtitle_font, fill="#66727F")
    draw.text((90, 205), "STATUS", font=font("regular", 19, True), fill="#1F4D78")
    legend_x = 205
    for status, status_color in STATUS_STYLES.items():
        pill_width = int(status_font.getlength(status)) + 36
        draw.rounded_rectangle(
            (legend_x, 190, legend_x + pill_width, 226),
            radius=18,
            fill=f"#{status_color}",
        )
        draw.text(
            (legend_x + pill_width / 2, 208),
            status,
            font=status_font,
            fill="#FFFFFF",
            anchor="mm",
        )
        legend_x += pill_width + 18
    row_top = 280
    row_height = 158
    row_gap = 80
    for index, (label, boxes, fill) in enumerate(layers):
        y1 = row_top + index * (row_height + row_gap)
        y2 = y1 + row_height
        draw.text(
            (90, y1 + row_height / 2),
            label.upper(),
            font=label_font,
            fill="#1F4D78",
            anchor="lm",
        )
        x_start = 385
        available = 1325
        gap = 28
        box_width = int((available - gap * (len(boxes) - 1)) / len(boxes))
        for box_index, (box, status) in enumerate(boxes):
            if status not in STATUS_STYLES:
                raise ValueError(f"Unsupported diagram status: {status}")
            x1 = x_start + box_index * (box_width + gap)
            x2 = x1 + box_width
            status_color = STATUS_STYLES[status]
            draw.rounded_rectangle(
                (x1, y1, x2, y2),
                radius=20,
                fill=fill,
                outline=f"#{status_color}",
                width=4,
            )
            pill_width = min(box_width - 28, int(status_font.getlength(status)) + 32)
            draw.rounded_rectangle(
                (x1 + 14, y1 + 12, x1 + 14 + pill_width, y1 + 45),
                radius=16,
                fill=f"#{status_color}",
            )
            draw.text(
                (x1 + 14 + pill_width / 2, y1 + 28),
                status,
                font=status_font,
                fill="#FFFFFF",
                anchor="mm",
            )
            wrapped(
                draw, box, (x1 + 18, y1 + 50, x2 - 18, y2 - 12), box_font, "#17324D"
            )
        if index < len(layers) - 1:
            next_y = row_top + (index + 1) * (row_height + row_gap)
            down_label, up_label = edge_labels[index]
            edge_mid = (y2 + next_y) / 2
            arrow(draw, (835, y2 + 8), (835, next_y - 8), "#6888A5")
            arrow(draw, (1210, next_y - 8), (1210, y2 + 8), "#6888A5")
            draw.text(
                (810, edge_mid), down_label, font=edge_font, fill="#4D6A82", anchor="rm"
            )
            draw.text(
                (1235, edge_mid), up_label, font=edge_font, fill="#4D6A82", anchor="lm"
            )
    draw.text(
        (90, image_height - 60),
        "IMPLEMENTED means repository code/test evidence; it never means a trained or deployed model.",
        font=font("regular", 20, True),
        fill="#7A2E28",
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, dpi=(180, 180))


def state_diagram(path: Path) -> None:
    image = Image.new("RGB", (1800, 900), "#F7F9FB")
    draw = ImageDraw.Draw(image)
    draw.text(
        (90, 68),
        "One run, explicit states",
        font=font("regular", 52, True),
        fill="#17324D",
    )
    draw.text(
        (90, 136),
        "Resumability and human release are part of the protocol, not UI decoration.",
        font=font("regular", 25),
        fill="#66727F",
    )
    nodes = [
        ("Created", f"#{PALE_BLUE}"),
        ("Running", f"#{PALE_BLUE}"),
        ("Awaiting approval", f"#{PALE_GOLD}"),
        ("Released", f"#{PALE_GREEN}"),
        ("Cancelled", f"#{PALE_RED}"),
        ("Failed", f"#{PALE_RED}"),
    ]
    positions = [
        (110, 290),
        (400, 290),
        (690, 290),
        (1050, 290),
        (690, 610),
        (400, 610),
    ]
    width, height = 235, 120
    for (label, fill), (x, y) in zip(nodes, positions):
        draw.rounded_rectangle(
            (x, y, x + width, y + height),
            radius=20,
            fill=fill,
            outline="#2E74B5",
            width=3,
        )
        wrapped(
            draw,
            label,
            (x + 12, y + 10, x + width - 12, y + height - 10),
            font("regular", 24, True),
            "#17324D",
        )
    arrow(draw, (345, 350), (392, 350))
    arrow(draw, (635, 350), (682, 350))
    arrow(draw, (925, 350), (1042, 350))
    arrow(draw, (807, 418), (807, 602), "#A66B00")
    arrow(draw, (690, 670), (643, 670), "#A33B30")
    arrow(draw, (517, 602), (517, 418), "#6888A5")
    transition_font = font("regular", 18, True)
    draw.text((368, 325), "start", font=transition_font, fill="#4D6A82", anchor="ms")
    draw.text(
        (658, 325), "candidate ready", font=transition_font, fill="#4D6A82", anchor="ms"
    )
    draw.text(
        (985, 325), "human approve", font=transition_font, fill="#4D6A82", anchor="ms"
    )
    draw.text((836, 500), "cancel / policy stop", font=transition_font, fill="#7F5C13")
    draw.text(
        (666, 646), "cleanup fault", font=transition_font, fill="#7A2E28", anchor="rm"
    )
    draw.text((540, 500), "retry / resume", font=transition_font, fill="#4D6A82")
    draw.text(
        (90, 805),
        "Every transition is persisted with optimistic concurrency and hash-chained events.",
        font=font("regular", 24, True),
        fill="#17324D",
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, dpi=(180, 180))


def kimi_adoption_diagram(path: Path) -> None:
    """Draw the evidence-gated Kimi K3 adoption map used by the handbook."""
    image = Image.new("RGB", (1800, 1250), "#F7F9FB")
    draw = ImageDraw.Draw(image)
    draw.text(
        (90, 62),
        "Kimi K3 to Hyperion: adoption decision map",
        font=font("regular", 48, True),
        fill="#17324D",
    )
    draw.text(
        (90, 128),
        "Paper claims pass an evidence-and-fit gate before they become target work.",
        font=font("regular", 24),
        fill="#66727F",
    )

    sources = [
        "Architecture\n§2–3",
        "Post-training\n§4",
        "Agent infrastructure\n§5",
        "Template + model card\nAppendix F",
    ]
    source_left = 110
    source_gap = 24
    source_width = 372
    source_top = 205
    source_bottom = 330
    for index, label in enumerate(sources):
        left = source_left + index * (source_width + source_gap)
        right = left + source_width
        draw.rounded_rectangle(
            (left, source_top, right, source_bottom),
            radius=18,
            fill="#E8EEF5",
            outline="#2E74B5",
            width=3,
        )
        wrapped(
            draw,
            label,
            (left + 18, source_top + 12, right - 18, source_bottom - 12),
            font("regular", 22, True),
            "#17324D",
        )
        centre = (left + right) // 2
        arrow(draw, (centre, source_bottom + 5), (centre, 370), "#6888A5")

    gate = (110, 380, 1690, 485)
    draw.rounded_rectangle(gate, radius=20, fill="#DDECF7", outline="#1F4D78", width=4)
    wrapped(
        draw,
        "HYPERION EVIDENCE + FEASIBILITY GATE — exact revision • rights • reproducibility • "
        "small-cluster fit • independent evaluation",
        (145, 392, 1655, 473),
        font("regular", 22, True),
        "#17324D",
    )

    decisions = [
        (
            "ADOPT IN TARGET",
            "#E8F4EC",
            [
                "Effort, token, time, and cost budgets",
                "Typed transcript and provider capability map",
                "Hidden verifier and submission limits",
                "Queue/KV admission and cache affinity",
            ],
        ),
        (
            "CONTROLLED EVALUATION",
            "#FFF4D6",
            [
                "Routed specialists before consolidation",
                "Structured judge scorepad and verbosity guard",
                "Resumable microVM rollouts",
                "QAT and acceptance-aware draft model",
            ],
        ),
        (
            "PRETRAINING RESEARCH",
            "#E8EEF5",
            [
                "KDA and Attention Residuals",
                "Stable LatentMoE, QB, and SiTU-GLU",
                "Per-head Muon and MoonEP",
                "Equal-data, equal-compute ablations",
            ],
        ),
    ]
    column_lefts = [110, 645, 1180]
    column_width = 510
    column_top = 580
    column_bottom = 1030
    bus_y = 535
    draw.line((900, 490, 900, bus_y), fill="#6888A5", width=6)
    draw.line((365, bus_y, 1435, bus_y), fill="#6888A5", width=6)
    for left, (heading, fill, items) in zip(column_lefts, decisions):
        right = left + column_width
        centre = (left + right) // 2
        arrow(draw, (centre, bus_y), (centre, column_top - 8), "#6888A5")
        draw.rounded_rectangle(
            (left, column_top, right, column_bottom),
            radius=22,
            fill=fill,
            outline="#2E74B5",
            width=3,
        )
        draw.rounded_rectangle(
            (left, column_top, right, column_top + 75),
            radius=22,
            fill="#1F4D78",
        )
        draw.rectangle((left, column_top + 50, right, column_top + 75), fill="#1F4D78")
        draw.text(
            (centre, column_top + 38),
            heading,
            font=font("regular", 22, True),
            fill="#FFFFFF",
            anchor="mm",
        )
        item_top = column_top + 92
        item_height = 80
        for item_index, item in enumerate(items):
            y1 = item_top + item_index * item_height
            draw.ellipse((left + 24, y1 + 25, left + 38, y1 + 39), fill="#2E74B5")
            wrapped(
                draw,
                item,
                (left + 52, y1, right - 22, y1 + 65),
                font("regular", 20, item_index == 0),
                "#17324D",
                spacing=5,
            )

    draw.rounded_rectangle(
        (110, 1080, 1690, 1188),
        radius=20,
        fill="#FBE9E7",
        outline="#A33B30",
        width=4,
    )
    wrapped(
        draw,
        "NOT A DIRECT RETROFIT — QLoRA cannot add KDA, AttnRes, or LatentMoE; the full "
        "2.78T / 104.2B-active K3 is not the small-cluster student; source claims are not "
        "Hyperion results.",
        (145, 1090, 1655, 1178),
        font("regular", 22, True),
        "#7A2E28",
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, dpi=(180, 180))


def build_diagrams(asset_dir: Path) -> None:
    layered_diagram(
        asset_dir / "hyperion-system-architecture.png",
        "Hyperion system architecture",
        "Authority moves downward through immutable records; evidence moves upward for review.",
        [
            (
                "Clients",
                [
                    ("Operator CLI", "IMPLEMENTED"),
                    ("Hosted API + worker", "IMPLEMENTED"),
                    ("Terminal adapter", "IMPLEMENTED"),
                ],
                "#E8EEF5",
            ),
            (
                "Coordination",
                [
                    ("Five-role workflow", "IMPLEMENTED"),
                    ("Budgets + approvals", "IMPLEMENTED"),
                    ("SQLite + artifact store", "IMPLEMENTED"),
                ],
                "#DDECF7",
            ),
            (
                "Model evidence",
                [
                    ("Corpus + teacher lineage\n(contracts only)", "IMPLEMENTED"),
                    ("Training + evaluation\n(harness only)", "IMPLEMENTED"),
                    ("Promotion + readiness\n(contracts only)", "IMPLEMENTED"),
                ],
                "#E8F4EC",
            ),
            (
                "Safety",
                [
                    ("Signed verifier authority", "IMPLEMENTED"),
                    ("Container controller", "IMPLEMENTED"),
                    ("Narrow .NET worker", "IMPLEMENTED"),
                ],
                "#FFF4D6",
            ),
        ],
        [
            ("request + approval", "quote + inert artifact"),
            ("frozen spec + inputs", "results + usage"),
            ("bound verification request", "verifier evidence"),
        ],
    )
    state_diagram(asset_dir / "hyperion-run-state-machine.png")
    layered_diagram(
        asset_dir / "hyperion-evidence-lifecycle.png",
        "Evidence lifecycle",
        "No gate inherits authority from an earlier gate without a new, bound record.",
        [
            (
                "Data",
                [
                    ("Rights + provenance tooling", "IMPLEMENTED"),
                    ("Production independent review", "EXTERNAL / BLOCKED"),
                    ("Signed corpus release\n(contract only)", "IMPLEMENTED"),
                ],
                "#E8EEF5",
            ),
            (
                "Training",
                [
                    ("Pinned base + runtime", "IMPLEMENTED"),
                    ("Real adapter checkpoint", "EXTERNAL / BLOCKED"),
                    ("Sealed artifact\n(contract only)", "IMPLEMENTED"),
                ],
                "#DDECF7",
            ),
            (
                "Evaluation",
                [
                    ("Production sealed tasks", "EXTERNAL / BLOCKED"),
                    ("Independent replay harness", "IMPLEMENTED"),
                    ("Promotion decision\n(contract only)", "IMPLEMENTED"),
                ],
                "#FFF4D6",
            ),
            (
                "Serving",
                [
                    ("Activation envelope", "IMPLEMENTED"),
                    ("Runtime attestation", "TARGET"),
                    ("Readiness proof\n(contract only)", "IMPLEMENTED"),
                ],
                "#E8F4EC",
            ),
        ],
        [
            ("approved release", "review findings"),
            ("candidate artifact", "accept / reject evidence"),
            ("promotion envelope", "readiness telemetry"),
        ],
    )
    layered_diagram(
        asset_dir / "hyperion-trust-chain.png",
        "Trust and signing chain",
        "Detached signers approve hashes and claims; they do not rebuild evidence supplied by the caller.",
        [
            (
                "Producers",
                [
                    ("Production corpus author", "EXTERNAL / BLOCKED"),
                    ("Trainer harness", "IMPLEMENTED"),
                    ("Verifier controller", "IMPLEMENTED"),
                ],
                "#E8EEF5",
            ),
            (
                "Independent review",
                [
                    ("Corpus reviewer", "EXTERNAL / BLOCKED"),
                    ("Replay authority", "EXTERNAL / BLOCKED"),
                    ("Promotion authority", "EXTERNAL / BLOCKED"),
                ],
                "#FFF4D6",
            ),
            (
                "Signing",
                [
                    ("KMS/HSM request path", "TARGET"),
                    ("Production detached signature", "EXTERNAL / BLOCKED"),
                    ("Canonical envelope", "IMPLEMENTED"),
                ],
                "#DDECF7",
            ),
            (
                "Consumers",
                [
                    ("Coordinator", "IMPLEMENTED"),
                    ("Hosted platform", "IMPLEMENTED"),
                    ("Deployment gate", "IMPLEMENTED"),
                ],
                "#E8F4EC",
            ),
        ],
        [
            ("evidence bundle", "findings + decision"),
            ("canonical preimage", "signature + key identity"),
            ("signed envelope", "verification result"),
        ],
    )
    layered_diagram(
        asset_dir / "hyperion-gpu-cluster-topology.png",
        "Small-cluster training topology",
        "Start with adapter training on a right-sized open-weight base; keep control and evidence off GPU workers.",
        [
            (
                "Control",
                [
                    ("Cluster scheduler + run registry", "TARGET"),
                    ("External secrets/KMS broker", "EXTERNAL / BLOCKED"),
                    ("Cluster metrics + audit", "TARGET"),
                ],
                "#E8EEF5",
            ),
            (
                "Storage",
                [
                    ("Production versioned corpus", "EXTERNAL / BLOCKED"),
                    ("Immutable artifact contracts", "IMPLEMENTED"),
                    ("Model/cache mirror", "TARGET"),
                ],
                "#DDECF7",
            ),
            (
                "GPU pool",
                [
                    ("Trainer rank 0", "EXTERNAL / BLOCKED"),
                    ("Trainer ranks 1..N", "EXTERNAL / BLOCKED"),
                    ("Evaluation / serving pool", "EXTERNAL / BLOCKED"),
                ],
                "#E8F4EC",
            ),
            (
                "Isolation",
                [
                    ("Verifier microVM", "TARGET"),
                    ("Restricted egress", "TARGET"),
                    ("Backup + restore", "TARGET"),
                ],
                "#FFF4D6",
            ),
        ],
        [
            ("signed run spec", "status + audit events"),
            ("read-only snapshots", "checkpoints + metrics"),
            ("verification request", "verdict + recovery proof"),
        ],
    )
    layered_diagram(
        asset_dir / "hyperion-frontier-model-factory.png",
        "Hyperion frontier model factory",
        "Models propose; governed evidence, independent verification, and human release grant authority.",
        [
            (
                "Authority",
                [
                    ("Registry + policy", "IMPLEMENTED"),
                    ("Budgets + credits", "IMPLEMENTED"),
                    ("Promotion + rollback\n(contracts only)", "IMPLEMENTED"),
                ],
                "#E8EEF5",
            ),
            (
                "Knowledge",
                [
                    ("Production reviewed corpus", "EXTERNAL / BLOCKED"),
                    ("Versioned private RAG", "IMPLEMENTED"),
                    ("Production sealed evaluation", "EXTERNAL / BLOCKED"),
                ],
                "#DDECF7",
            ),
            (
                "Teachers",
                [
                    ("Hosted frontier teacher", "TARGET"),
                    ("Open large models", "TARGET"),
                    ("Critics + judges", "TARGET"),
                ],
                "#FFF4D6",
            ),
            (
                "Student factory",
                [
                    ("Raw-base bake-off\n(synthetic evidence)", "IMPLEMENTED"),
                    ("SFT / DPO / RLVR\n(harness only)", "IMPLEMENTED"),
                    ("Effort specialists + router", "RESEARCH"),
                ],
                "#E8F4EC",
            ),
            (
                "Runtime",
                [
                    ("Coordinator + tools", "IMPLEMENTED"),
                    ("Verifier + replay", "IMPLEMENTED"),
                    ("Inference + GPU cluster", "EXTERNAL / BLOCKED"),
                ],
                "#E8EEF5",
            ),
        ],
        [
            ("policy + rights gates", "evidence + exceptions"),
            ("authorised examples", "teacher traces + critiques"),
            ("training specification", "candidate + metrics"),
            ("accepted artifact", "telemetry + rollback evidence"),
        ],
    )
    kimi_adoption_diagram(asset_dir / "hyperion-kimi-k3-adoption-map.png")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, column_count: int) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    tbl_pr.append(indent)
    margins = OxmlElement("w:tblCellMar")
    for side, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    tbl_pr.append(margins)
    base = CONTENT_DXA // column_count
    widths = [base] * column_count
    widths[-1] += CONTENT_DXA - sum(widths)
    for row in table.rows:
        prevent_row_split(row)
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([color, underline])
    run.append(props)
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|https?://[^\s]+|\*[^*]+\*)")


def soft_wrap(value: str) -> str:
    if len(value) < 28 or " " in value:
        return value
    value = re.sub(r"([/\\._-])", r"\1" + "\u200b", value)
    return re.sub(r"([A-Za-z0-9]{16})(?=[A-Za-z0-9])", r"\1" + "\u200b", value)


def add_inline(paragraph, value: str, *, base_bold: bool = False) -> None:
    position = 0
    for match in INLINE.finditer(value):
        if match.start() > position:
            paragraph.add_run(value[position:match.start()]).bold = base_bold
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(soft_wrap(token[1:-1]))
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
        elif token.startswith("["):
            link_match = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            assert link_match
            add_hyperlink(paragraph, link_match.group(1), link_match.group(2))
        elif token.startswith("http"):
            url = token.rstrip(".,;")
            parsed = urlparse(url)
            path_parts = [unquote(part) for part in parsed.path.split("/") if part]
            path_label = "/".join(path_parts[-2:]) if path_parts else "source"
            link_label = f"{parsed.netloc.removeprefix('www.')} / {path_label}"
            add_hyperlink(paragraph, link_label, url)
            if len(url) < len(token):
                paragraph.add_run(token[len(url):])
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        position = match.end()
    if position < len(value):
        paragraph.add_run(value[position:]).bold = base_bold


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string("25313B")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.widow_control = True
    for name in ("List Bullet", "List Bullet 2", "List Bullet 3", "List Number", "List Number 2", "List Number 3"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25
        level = 0 if name.endswith(("Bullet", "Number")) else int(name[-1]) - 1
        style.paragraph_format.left_indent = Inches(0.375 + 0.25 * level)
        style.paragraph_format.first_line_indent = Inches(-0.188)
    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code Block"]
    code.font.name = "Consolas"
    code.font.size = Pt(8.3)
    code.font.color.rgb = RGBColor.from_string("243B53")
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.18)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(7)
    if "Figure Caption" not in styles:
        caption = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = styles["Figure Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(GREY)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = False
    caption.paragraph_format.keep_together = True


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True
    header = section.header.paragraphs[0]
    header.text = "HYPERION  /  ENGINEERING HANDBOOK"
    header.style = document.styles["Normal"]
    header.runs[0].font.size = Pt(8.5)
    header.runs[0].font.bold = True
    header.runs[0].font.color.rgb = RGBColor.from_string(DARK_BLUE)
    footer = section.footer.paragraphs[0]
    footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    run = footer.add_run("PRIVATE ENGINEERING MATERIAL")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GREY)
    footer.add_run("\t")
    add_page_field(footer)
    for footer_run in footer.runs:
        footer_run.font.size = Pt(8)
        footer_run.font.color.rgb = RGBColor.from_string(GREY)


def configure_default_language(document: Document, language: str) -> None:
    styles = document.styles.element
    doc_defaults = styles.find(qn("w:docDefaults"))
    if doc_defaults is None:
        doc_defaults = OxmlElement("w:docDefaults")
        styles.insert(0, doc_defaults)
    run_properties_default = doc_defaults.find(qn("w:rPrDefault"))
    if run_properties_default is None:
        run_properties_default = OxmlElement("w:rPrDefault")
        doc_defaults.insert(0, run_properties_default)
    run_properties = run_properties_default.find(qn("w:rPr"))
    if run_properties is None:
        run_properties = OxmlElement("w:rPr")
        run_properties_default.append(run_properties)
    language_element = run_properties.find(qn("w:lang"))
    if language_element is None:
        language_element = OxmlElement("w:lang")
        run_properties.append(language_element)
    language_element.set(qn("w:val"), language)
    language_element.set(qn("w:eastAsia"), language)
    language_element.set(qn("w:bidi"), language)

    normal_run_properties = document.styles["Normal"].element.get_or_add_rPr()
    normal_language = normal_run_properties.find(qn("w:lang"))
    if normal_language is None:
        normal_language = OxmlElement("w:lang")
        normal_run_properties.append(normal_language)
    normal_language.set(qn("w:val"), language)
    normal_language.set(qn("w:eastAsia"), language)
    normal_language.set(qn("w:bidi"), language)


def configure_extended_properties(document: Document) -> None:
    text_chunks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text_chunks.extend(paragraph.text for paragraph in cell.paragraphs)
    plain_text = "\n".join(text_chunks)
    word_count = len(re.findall(r"\b[\w]+(?:[-'][\w]+)*\b", plain_text))
    properties = {
        "Template": "Hyperion Engineering Handbook",
        "Pages": str(HANDBOOK_EXPECTED_PAGES),
        "Words": str(word_count),
        "Characters": str(sum(not character.isspace() for character in plain_text)),
        "Application": "Hyperion Handbook Builder",
        "Paragraphs": str(len(text_chunks)),
        "CharactersWithSpaces": str(len(plain_text)),
        "AppVersion": HANDBOOK_VERSION,
    }
    app_part = next(
        part
        for part in document.part.package.parts
        if str(part.partname) == "/docProps/app.xml"
    )
    root = etree.fromstring(app_part.blob)
    namespace = (
        "{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}"
    )
    for name, value in properties.items():
        element = root.find(f"{namespace}{name}")
        if element is None:
            element = etree.SubElement(root, f"{namespace}{name}")
        element.text = value
    app_part._blob = etree.tostring(
        root,
        encoding="UTF-8",
        xml_declaration=True,
        standalone=True,
    )


def add_cover(document: Document) -> None:
    document.add_paragraph().paragraph_format.space_after = Pt(38)
    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run(f"PRIVATE ENGINEERING HANDBOOK  •  VERSION {HANDBOOK_VERSION}")
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(24)
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run("Hyperion")
    run.bold = True
    run.font.size = Pt(30)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    run = subtitle.add_run("LLM Coordinator\nEngineering Handbook")
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    band = document.add_paragraph()
    band.alignment = WD_ALIGN_PARAGRAPH.CENTER
    band.paragraph_format.left_indent = Inches(0.12)
    band.paragraph_format.right_indent = Inches(0.12)
    band.paragraph_format.space_before = Pt(14)
    band.paragraph_format.space_after = Pt(14)
    band_properties = band._p.get_or_add_pPr()
    band_shading = OxmlElement("w:shd")
    band_shading.set(qn("w:fill"), NAVY)
    band_properties.append(band_shading)
    run = band.add_run(
        "Coordinator • frontier model factory • isolated verifier\n"
        "For ML, platform, safety, and operations teams"
    )
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)
    document.add_paragraph().paragraph_format.space_after = Pt(54)
    for label, value in (
        ("DATE", HANDBOOK_DATE),
        ("REPOSITORY", "github.com/dhruuvsharma/Hyperion"),
        ("CLASSIFICATION", "Private engineering material"),
        (
            "STATUS",
            "Infrastructure implemented; no trained Hyperion checkpoint or production deployment",
        ),
    ):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(8)
        label_run = paragraph.add_run(f"{label}  ")
        label_run.bold = True
        label_run.font.size = Pt(9)
        label_run.font.color.rgb = RGBColor.from_string(BLUE)
        value_run = paragraph.add_run(value)
        value_run.font.size = Pt(10.5)
        value_run.font.color.rgb = RGBColor.from_string("25313B")
    document.add_page_break()


def add_callout(document: Document, value: str, fill: str = PALE_GOLD) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(9)
    paragraph.paragraph_format.keep_together = True
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), BLUE)
    borders.append(left)
    p_pr.append(borders)
    add_inline(paragraph, value)


def markdown_cells(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def add_table(document: Document, rows: list[list[str]]) -> None:
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    for row_index, values in enumerate(rows):
        for column_index in range(column_count):
            cell = table.cell(row_index, column_index)
            value = values[column_index] if column_index < len(values) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.08
            add_inline(paragraph, value, base_bold=row_index == 0)
            for run in paragraph.runs:
                run.font.size = Pt(8.3)
            if row_index == 0:
                set_cell_shading(cell, PALE_BLUE)
                cell._tc.get_or_add_tcPr().append(OxmlElement("w:tcFitText"))
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, column_count)
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_picture(
    document: Document, source_dir: Path, alt: str, relative_path: str, number: int
) -> None:
    image_path = (source_dir / relative_path).resolve()
    if not image_path.exists():
        add_callout(
            document, f"Diagram unavailable during build: {relative_path}", PALE_RED
        )
        return
    description = IMAGE_DESCRIPTIONS.get(relative_path, alt)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(6.35))
    drawing = run._r.xpath(".//wp:docPr")
    if drawing:
        drawing[0].set("descr", description)
        drawing[0].set("title", alt)
    caption = document.add_paragraph(style="Figure Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_together = True
    caption.add_run(f"Figure {number}. {alt}")


def add_code(document: Document, code: str) -> None:
    paragraph = document.add_paragraph(style="Code Block")
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F5F7")
    p_pr.append(shading)
    paragraph.paragraph_format.keep_together = len(code.splitlines()) <= 22
    paragraph.add_run(code.rstrip())


def create_numbering_instance(document: Document, start: int = 1) -> int:
    """Create a real, independently restarting decimal list."""
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(element.get(qn("w:abstractNumId")))
        for element in numbering.findall(qn("w:abstractNum"))
    ]
    number_ids = [
        int(element.get(qn("w:numId")))
        for element in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    number_id = max(number_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "multilevel")
    abstract.append(multi_level)
    for level in range(3):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        start_element = OxmlElement("w:start")
        start_element.set(qn("w:val"), "1")
        number_format = OxmlElement("w:numFmt")
        number_format.set(qn("w:val"), "decimal")
        level_text = OxmlElement("w:lvlText")
        level_text.set(qn("w:val"), f"%{level + 1}.")
        suffix = OxmlElement("w:suff")
        suffix.set(qn("w:val"), "space")
        justification = OxmlElement("w:lvlJc")
        justification.set(qn("w:val"), "left")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), str(540 + 360 * level))
        tabs.append(tab)
        indent = OxmlElement("w:ind")
        indent.set(qn("w:left"), str(540 + 360 * level))
        indent.set(qn("w:hanging"), "270")
        p_pr.extend([tabs, indent])
        lvl.extend([start_element, number_format, level_text, suffix, justification, p_pr])
        abstract.append(lvl)
    numbering.append(abstract)

    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(number_id))
    abstract_reference = OxmlElement("w:abstractNumId")
    abstract_reference.set(qn("w:val"), str(abstract_id))
    number.append(abstract_reference)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), str(start))
    override.append(start_override)
    number.append(override)
    numbering.append(number)
    return number_id


def apply_numbering(paragraph, number_id: int, level: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), str(number_id))
    num_pr.extend([ilvl, num_id])
    p_pr.append(num_pr)


def add_contents(document: Document, headings: list[str], source_dir: Path) -> None:
    heading = document.add_paragraph("Section index", style="Heading 1")
    heading.paragraph_format.space_before = Pt(0)
    document.add_paragraph(
        "This compact index lists the handbook's numbered sections; the PDF outline and Word headings provide "
        "navigation."
    )
    for value in headings:
        paragraph = document.add_paragraph(style="List Number")
        # Markdown section headings already carry explicit numbers; avoid a second visible number.
        paragraph.style = document.styles["Normal"]
        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        paragraph.paragraph_format.space_after = Pt(3)
        add_inline(paragraph, value)
    add_callout(
        document,
        "Evidence boundary: Hyperion has implemented and tested infrastructure, but no rights-approved "
        "production corpus, conforming training run, trained checkpoint, accepted benchmark, or production "
        "deployment. Every quality and readiness claim must remain attached to its evidence.",
    )
    add_picture(
        document,
        source_dir,
        "Hyperion system architecture",
        "assets/hyperion-system-architecture.png",
        1,
    )
    document.add_page_break()


def render_markdown(document: Document, source: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    headings = [
        line[3:].strip()
        for line in lines
        if line.startswith("## ") and not line.startswith("## Engineering")
    ]
    add_contents(document, headings, source.parent)
    start = next(
        index for index, line in enumerate(lines) if re.match(r"##\s+1\.", line)
    )
    index = start
    figure_number = 2
    active_number_id: int | None = None
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            active_number_id = None
            index += 1
            continue
        image_match = re.match(r"!\[([^]]+)\]\(([^)]+)\)", stripped)
        heading_match = re.match(r"^(#{2,4})\s+(.+)$", line)
        if heading_match:
            active_number_id = None
            level = len(heading_match.group(1)) - 1
            text = heading_match.group(2).strip()
            document.add_paragraph(text, style=f"Heading {min(level, 3)}")
            index += 1
            continue
        if image_match:
            active_number_id = None
            add_picture(
                document,
                source.parent,
                image_match.group(1),
                image_match.group(2),
                figure_number,
            )
            figure_number += 1
            index += 1
            continue
        if stripped.startswith("```"):
            active_number_id = None
            index += 1
            code_lines: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            add_code(document, "\n".join(code_lines))
            index += 1
            continue
        if stripped.startswith(">"):
            active_number_id = None
            quote: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote.append(lines[index].strip()[1:].strip())
                index += 1
            add_callout(document, " ".join(quote))
            continue
        if (
            stripped.startswith("|")
            and index + 1 < len(lines)
            and re.match(r"^\s*\|?\s*:?-{3,}", lines[index + 1])
        ):
            active_number_id = None
            rows = [markdown_cells(line)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(markdown_cells(lines[index]))
                index += 1
            add_table(document, rows)
            continue
        list_match = re.match(r"^(\s*)[-*]\s+(.+)$", line)
        number_match = re.match(r"^(\s*)(\d+)\.\s+(.+)$", line)
        if list_match or number_match:
            match = list_match or number_match
            assert match
            level = min(len(match.group(1).replace("\t", "    ")) // 2, 2)
            value = match.group(2) if list_match else match.group(3)
            index += 1
            while (
                index < len(lines)
                and lines[index].startswith("  ")
                and not re.match(r"^\s*[-*]\s+", lines[index])
            ):
                value += " " + lines[index].strip()
                index += 1
            if list_match:
                active_number_id = None
                style_root = "List Bullet"
                style = style_root if level == 0 else f"{style_root} {level + 1}"
            else:
                if active_number_id is None:
                    active_number_id = create_numbering_instance(
                        document, int(number_match.group(2))
                    )
                style = "List Number" if level == 0 else f"List Number {level + 1}"
            paragraph = document.add_paragraph(style=style)
            if number_match:
                apply_numbering(paragraph, active_number_id, level)
            add_inline(paragraph, value)
            continue
        if stripped in {"---", "***"}:
            active_number_id = None
            index += 1
            continue
        active_number_id = None
        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index]
            candidate_stripped = candidate.strip()
            if not candidate_stripped:
                break
            if (
                candidate_stripped.startswith(("#", ">", "```", "![", "|"))
                or re.match(r"^\s*[-*]\s+", candidate)
                or re.match(r"^\s*\d+\.\s+", candidate)
                or candidate_stripped in {"---", "***"}
            ):
                break
            paragraph_lines.append(candidate_stripped)
            index += 1
        paragraph = document.add_paragraph()
        add_inline(paragraph, " ".join(paragraph_lines))


def validate_source_metadata(source: Path) -> None:
    content = source.read_text(encoding="utf-8")
    required = [
        f"**Version:** {HANDBOOK_VERSION}",
        f"**Date:** {HANDBOOK_DATE}",
    ]
    missing = [marker for marker in required if marker not in content]
    if missing:
        raise ValueError(
            f"Handbook source metadata does not match builder constants: {missing}"
        )


def build(source: Path, output: Path) -> None:
    validate_source_metadata(source)
    assets = source.parent / "assets"
    build_diagrams(assets)
    document = Document()
    configure_styles(document)
    configure_document(document)
    configure_default_language(document, "en-GB")
    document.core_properties.title = "Hyperion LLM Coordinator Engineering Handbook"
    document.core_properties.subject = "Frontier model architecture, model evidence, verifier safety, and ML operations"
    document.core_properties.author = "Hyperion Engineering"
    document.core_properties.last_modified_by = "Hyperion Engineering"
    document.core_properties.keywords = "Hyperion, LLM coordinator, distillation, RLVR, QLoRA, verifier, model governance"
    document.core_properties.comments = "Private engineering material"
    document.core_properties.category = "Engineering handbook"
    document.core_properties.content_status = "Reviewed"
    document.core_properties.created = HANDBOOK_RELEASED_UTC
    document.core_properties.modified = HANDBOOK_RELEASED_UTC
    document.core_properties.revision = 1
    document.core_properties.version = HANDBOOK_VERSION
    settings = document.settings.element
    language = settings.find(qn("w:themeFontLang"))
    if language is None:
        language = OxmlElement("w:themeFontLang")
        settings.append(language)
    language.set(qn("w:val"), "en-GB")
    language.set(qn("w:eastAsia"), "en-GB")
    language.set(qn("w:bidi"), "en-GB")
    add_cover(document)
    render_markdown(document, source)
    configure_extended_properties(document)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=Path("docs/engineering/hyperion-llm-coordinator-handbook.md"))
    parser.add_argument("--output", type=Path, default=Path("docs/engineering/hyperion-llm-coordinator-handbook.docx"))
    arguments = parser.parse_args()
    build(arguments.source.resolve(), arguments.output.resolve())
    print(arguments.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
